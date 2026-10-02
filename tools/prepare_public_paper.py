"""Create the v0.2.2 public paper copy from the retained source document.

The source document is never modified. This script applies only the disclosed
publication corrections, preserves the existing figures, and repeats the first
row of every table for reader continuity when a table spans a page.
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
# The retained source lives beside this repository during the local release
# build.  Resolve it from the repository layout rather than embedding a
# machine-specific path in a public release helper.
SOURCE_DOCUMENT = REPOSITORY_ROOT.parent / "White语义白盒治理.docx"
OUTPUT_DOCUMENT = REPOSITORY_ROOT / "paper" / "White语义白盒治理_v0.2.2.docx"


def set_paragraph_text(paragraph, expected: str, replacement: str) -> None:
    if paragraph.text != expected:
        raise ValueError(f"Expected paragraph not found: {expected!r}")
    paragraph.text = replacement


def repeat_header_row(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def clear_numbering(paragraph) -> None:
    properties = paragraph._p.pPr
    if properties is not None and properties.numPr is not None:
        properties.remove(properties.numPr)


def copy_numbering(source, target) -> None:
    source_properties = source._p.pPr
    if source_properties is None or source_properties.numPr is None:
        raise ValueError("Expected numbered source paragraph")
    clear_numbering(target)
    target._p.get_or_add_pPr().append(deepcopy(source_properties.numPr))


def replace_table_rows(table, rows) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    header = table.rows[0]
    header.cells[0].text = "Item"
    header.cells[1].text = "Release status"
    for item, value in rows:
        cells = table.add_row().cells
        cells[0].text = item
        cells[1].text = value


def compact_related_work_table(table) -> None:
    """Make the five-column related-work comparison legible on A4 portrait.

    The source table has a useful distinction but its five narrow columns force
    single words onto separate lines after PDF export.  Preserve every claim,
    while grouping the contextual columns into a compact three-column layout.
    """
    source_rows = [tuple(cell.text for cell in row.cells) for row in table.rows]
    for row, values in zip(table.rows, source_rows):
        category, representative, problem, non_replacement, distinction = values
        row.cells[0].text = category
        scope_cell = row.cells[1].merge(row.cells[2])
        scope_cell.text = (
            f"{representative}\n"
            f"Scope: {problem}\n"
            f"Not claimed to replace: {non_replacement}"
        )
        distinction_cell = row.cells[3].merge(row.cells[4])
        distinction_cell.text = distinction
    header = table.rows[0]
    header.cells[0].text = "Category"
    header.cells[1].text = "Representative work and scope"
    header.cells[3].text = "Narrow distinction studied here"


def main() -> None:
    if not SOURCE_DOCUMENT.is_file():
        raise SystemExit(f"Source document not found: {SOURCE_DOCUMENT}")
    OUTPUT_DOCUMENT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE_DOCUMENT, OUTPUT_DOCUMENT)
    document = Document(OUTPUT_DOCUMENT)

    paragraphs = document.paragraphs
    set_paragraph_text(
        paragraphs[3],
        "Technical Research Note v0.2 — Draft for Public Review",
        "Technical Research Note v0.2.2 — Public Release Candidate",
    )
    set_paragraph_text(
        paragraphs[4],
        "Author: [To be confirmed]\nAffiliation: [Independent Researcher / To be confirmed]\nDate: October 2026",
        "Author: 万根\nAffiliation: Independent Researcher\nCorrespondence: 593315006@qq.com\nDate: October 2026",
    )
    paragraphs[9].text = paragraphs[9].text.replace(
        "passed 10 tests under Python 3.9 and 10 tests under Python 3.13.",
        "passed 14 tests under Python 3.9 and 14 tests under Python 3.13.",
    )
    set_paragraph_text(
        paragraphs[142],
        "The implementation fingerprints the evidence representation used for the decision. A replay therefore cannot silently substitute a different evidence object while claiming that the original decision is being reproduced.",
        "The implementation fingerprints the evidence representation used for the decision. An idempotency replay is accepted only when the candidate, evidence, policy snapshot, and original state snapshot all match the first submission. A changed bound input is rejected rather than silently returning the earlier ALLOW record.",
    )
    set_paragraph_text(paragraphs[186], "APPLIED or FAILED", "APPLIED, FAILED, or UNKNOWN")
    set_paragraph_text(paragraphs[193], "NOT_RUN", "UNKNOWN\nNOT_RUN")
    set_paragraph_text(paragraphs[235], "NOT_RUN → do not advance", "UNKNOWN → do not advance\nNOT_RUN → do not advance")
    set_paragraph_text(
        paragraphs[545],
        "[9] T. Shi et al., “Progent: Programmable Privilege Control for LLM Agents,” arXiv:2504.11703, 2025. https://arxiv.org/abs/2504.11703.",
        "[9] T. Shi et al., “Progent: Securing AI Agents with Privilege Control,” arXiv:2504.11703, 2025. https://arxiv.org/abs/2504.11703.",
    )
    set_paragraph_text(paragraphs[578], "    status: APPLIED | FAILED | NOT_RUN", "    status: APPLIED | FAILED | UNKNOWN | NOT_RUN")
    set_paragraph_text(paragraphs[615], "Appendix E. Information Requiring Author Confirmation", "Appendix E. Release Metadata")

    # The source document left list-numbering properties on three empty
    # paragraphs, which rendered as stray 1/2/3 markers before the two actual
    # contribution statements. Keep the original prose but attach the list
    # numbering only to the two populated contribution paragraphs.
    copy_numbering(paragraphs[61], paragraphs[62])
    copy_numbering(paragraphs[61], paragraphs[65])
    for paragraph in (paragraphs[61], paragraphs[63], paragraphs[64], paragraphs[66]):
        clear_numbering(paragraph)

    table2 = document.tables[2]
    cells = table2.add_row().cells
    cells[0].text = "Idempotency Bound Input Conflict"
    cells[1].text = "Candidate, evidence, policy, or original state differs under the same idempotency key"
    cells[2].text = "DENY; no second effect"
    cells[3].text = "Demonstrated"

    table6 = document.tables[6]
    for row in table6.rows:
        if row.cells[0].text == "Results reproduce across Python 3.9 and 3.13":
            row.cells[1].text = "14 tests on each version"
    cells = table6.add_row().cells
    cells[0].text = "Changed bound input under same idempotency key is rejected"
    cells[1].text = "Negative subcases in replay test"
    cells[2].text = "DEMONSTRATED_IN_REFERENCE_IMPLEMENTATION"

    replace_table_rows(
        document.tables[7],
        [
            ("Author", "万根"),
            ("Affiliation", "Independent Researcher"),
            ("Release version", "0.2.2"),
            ("Citation metadata", "CITATION.cff"),
            ("Code and paper license", "LICENSE.md"),
            ("Public contact channel", "593315006@qq.com"),
            ("Public GitHub repository", "Published separately with this release"),
            ("Git tag", "v0.2.2-reference-demo after contact-update freeze"),
            ("Zenodo record and DOI", "Not yet created"),
            ("Release manifest and checksums", "Generated in this release package"),
            ("Protected-capability overview", "Included without protected implementation detail"),
            ("Patent and public-disclosure boundary", "Requires independent legal review before publication"),
        ],
    )

    compact_related_work_table(document.tables[4])

    for table in document.tables:
        repeat_header_row(table.rows[0])

    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = "White Box State Transition Governance"
    document.core_properties.subject = "Public reference release v0.2.2"
    document.core_properties.comments = ""
    document.save(OUTPUT_DOCUMENT)
    print(OUTPUT_DOCUMENT)


if __name__ == "__main__":
    main()
